"""Outreach log (Excel) and do-not-contact list."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HEADERS = ["Date", "Name", "Company", "Email", "Phone", "Channel", "Status", "Notes"]
WIDTHS = [20, 24, 24, 32, 18, 11, 12, 48]


class OutreachLog:
    def __init__(self, path: Path):
        self.path = path

    def _open(self) -> tuple[Any, Any]:
        if self.path.exists():
            wb = load_workbook(self.path)
            return wb, wb.active
        wb = Workbook()
        ws = wb.active
        ws.title = "Outreach"
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for i, w in enumerate(WIDTHS):
            ws.column_dimensions[chr(ord("A") + i)].width = w
        ws.freeze_panes = "A2"
        return wb, ws

    def add(
        self,
        name: str,
        email: str = "",
        phone: str = "",
        channel: str = "",
        status: str = "contacted",
        company: str = "",
        notes: str = "",
        when: datetime | None = None,
    ) -> dict[str, str]:
        if not name.strip():
            raise ValueError("A contact name is required.")
        row = [
            (when or datetime.now()).strftime("%Y-%m-%d %H:%M"),
            name.strip(),
            company.strip(),
            email.strip(),
            phone.strip(),
            channel,
            status,
            notes,
        ]
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb, ws = self._open()
        ws.append(row)
        wb.save(self.path)
        wb.close()
        return dict(zip(HEADERS, row, strict=False))

    def rows(self) -> list[dict[str, str]]:
        if not self.path.exists():
            return []
        wb = load_workbook(self.path, read_only=True)
        try:
            data = list(wb.active.iter_rows(values_only=True))
        finally:
            wb.close()
        return [
            {h: ("" if v is None else str(v)) for h, v in zip(HEADERS, r, strict=False)}
            for r in data[1:]
            if any(r)
        ]


class DoNotContact:
    """People who asked not to be contacted. Checked before every send."""

    def __init__(self, path: Path):
        self.path = path

    def _load(self) -> set[str]:
        try:
            return set(json.loads(self.path.read_text("utf-8")))
        except (OSError, ValueError):
            return set()

    @staticmethod
    def _key(value: str) -> str:
        v = value.strip().lower()
        return "".join(ch for ch in v if ch.isdigit()) if "@" not in v else v

    def add(self, value: str) -> None:
        items = self._load()
        items.add(self._key(value))
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(sorted(items), indent=1), "utf-8")

    def blocked(self, value: str) -> bool:
        return bool(value.strip()) and self._key(value) in self._load()
